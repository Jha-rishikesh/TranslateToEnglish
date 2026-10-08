import os
os.environ["ARGOS_CHUNK_TYPE"] = "MINISBD"

from docx import Document
import argostranslate.translate


def should_translate(text):
    text = text.strip()

    if not text:
        return False

    if len(text) <= 1:
        return False

    if text.isdigit():
        return False

    japanese_markers = [
        "Ō", "ō", "chō",
        "shi", "ku",
        "aza", "miya"
    ]

    for marker in japanese_markers:
        if marker in text:
            return False

    return True


def translate_docx(
    input_file,
    output_file,
    source_language,
    target_language="en"
):
    installed_languages = (
        argostranslate.translate.get_installed_languages()
    )

    source = next(
        lang for lang in installed_languages
        if lang.code == source_language
    )

    target = next(
        lang for lang in installed_languages
        if lang.code == target_language
    )

    translator = source.get_translation(target)

    doc = Document(input_file)

    # Translate normal paragraphs
    for paragraph in doc.paragraphs:
        original_text = paragraph.text

        if not should_translate(original_text):
            continue

        translated_text = translator.translate(original_text)

        if paragraph.runs:
            paragraph.runs[0].text = translated_text

            for run in paragraph.runs[1:]:
                run.text = ""

        print("Original:", original_text)
        print("English :", translated_text)
        print("-" * 40)

    # Translate table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                original_text = cell.text

                if not should_translate(original_text):
                    continue

                translated_text = translator.translate(original_text)

                # Preserve formatting using paragraph runs
                for paragraph in cell.paragraphs:
                    if not paragraph.text.strip():
                        continue

                    if paragraph.runs:
                        paragraph.runs[0].text = translated_text

                        for run in paragraph.runs[1:]:
                            run.text = ""

                    break

                print("Table:", original_text, "->", translated_text)

    doc.save(output_file)

    print("\nDOCX translation completed")
    print("Output:", output_file)

    if __name__ == "__main__":

      translate_docx(
          "input/DummyFrench.docx",
          "output/DummyFrench_English.docx",
          "fr"
      )
