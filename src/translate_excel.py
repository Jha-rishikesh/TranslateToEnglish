import os
os.environ["ARGOS_CHUNK_TYPE"] = "MINISBD"

from openpyxl import load_workbook
import argostranslate.translate


def should_translate(text):
    text = text.strip()

    if not text:
        return False

    if len(text) <= 1:
        return False

    if text.isdigit():
        return False

    japanese_markers = [
        "Ō", "ō", "chō",
        "shi", "ku",
        "aza", "miya"
    ]

    for marker in japanese_markers:
        if marker in text:
            return False

    return True


def translate_excel(
    input_file,
    output_file,
    source_language,
    target_language="en"
):
    installed_languages = (
        argostranslate.translate.get_installed_languages()
    )

    source = next(
        lang for lang in installed_languages
        if lang.code == source_language
    )

    target = next(
        lang for lang in installed_languages
        if lang.code == target_language
    )

    translator = source.get_translation(target)

    workbook = load_workbook(input_file)

    for sheet in workbook.worksheets:
        print("\nSheet:", sheet.title)

        for row in sheet.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue

                original_text = cell.value

                if not should_translate(original_text):
                    continue

                translated_text = translator.translate(original_text)
                cell.value = translated_text

                print(
                    cell.coordinate,
                    ":",
                    original_text,
                    "->",
                    translated_text
                )

    workbook.save(output_file)

    print("\nExcel translation completed")
    print("Output:", output_file)

if __name__ == "__main__":
    
    translate_excel(
      "input/DummyFrench.xlsx",
      "output/DummyFrench_English.xlsx",
      "fr"
  )    