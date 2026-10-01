from docx import Document

doc = Document("input/DummyFrench.docx")

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Total Sections: {len(doc.sections)}")