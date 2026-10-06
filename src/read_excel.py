from openpyxl import load_workbook

input_file = "input/DummyFrench.xlsx"

workbook = load_workbook(input_file)

print("Sheets:", workbook.sheetnames)

for sheet in workbook.worksheets:
    print("\nSheet:", sheet.title)

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                print(cell.coordinate, "=", cell.value)