from pptx import Presentation

ppt = Presentation("input/DummyFrench.pptx")

for slide_no, slide in enumerate(ppt.slides, start=1):

    print(f"\n=== Slide {slide_no} ===")

    for shape in slide.shapes:

        if hasattr(shape, "text"):

            if shape.text.strip():

                print(shape.text)