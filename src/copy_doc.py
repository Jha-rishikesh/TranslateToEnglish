from docx import Document

source_doc = Document("input/DummyFrench.docx")

new_doc = Document()

for para in source_doc.paragraphs:
    new_para = new_doc.add_paragraph(para.text)
    new_para.style = para.style

new_doc.save("output/DummyFrench_Copy.docx")

print("File Created Successfully")