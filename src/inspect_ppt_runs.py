from pptx import Presentation

ppt = Presentation("input/DummyFrench.pptx")

for slide_no, slide in enumerate(ppt.slides, start=1):

    print(f"\n{'=' * 50}")
    print(f"SLIDE {slide_no}")

    for shape_no, shape in enumerate(slide.shapes, start=1):

        if not hasattr(shape, "text_frame"):
            continue

        print(f"\nShape {shape_no}")

        for para_no, paragraph in enumerate(shape.text_frame.paragraphs, start=1):

            print(f"\nParagraph {para_no}")

            for run_no, run in enumerate(paragraph.runs, start=1):

                print(
                    f"Run {run_no}: "
                    f"'{run.text}'"
                )
