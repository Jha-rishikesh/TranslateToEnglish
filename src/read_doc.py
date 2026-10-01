from docx import Document

doc = Document("input/DummyFrench.docx")

# for para in doc.paragraphs:
#     if para.text.strip():
#         print(para.text)
###############
for i, para in enumerate(doc.paragraphs, start=1):
  print(f"\nParagraph {i}")
  print(f"Style: {para.style.name}")
  print(f"Text : {para.text}")