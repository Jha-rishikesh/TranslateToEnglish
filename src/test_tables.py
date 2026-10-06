from docx import Document

doc = Document("input/DummyFrench.docx")

print("Tables Found:", len(doc.tables))

for table_no, table in enumerate(doc.tables, start=1):

    print(f"\nTable {table_no}")

    for row in table.rows:
        for cell in row.cells:
            print(cell.text)