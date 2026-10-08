import os
os.environ["ARGOS_CHUNK_TYPE"] = "MINISBD"

from pptx import Presentation
import argostranslate.translate


def should_translate(text):
    text = text.strip()

    if not text:
        return False

    if len(text) <= 1:
        return False

    if text.isdigit():
        return False

    japanese_markers = [
        "Ō", "ō", "chō",
        "shi", "ku",
        "aza", "miya"
    ]

    for marker in japanese_markers:
        if marker in text:
            return False

    return True


def translate_ppt(
    input_file,
    output_file,
    source_language,
    target_language="en"
):
    installed_languages = (
        argostranslate.translate.get_installed_languages()
    )

    source = next(
        lang for lang in installed_languages
        if lang.code == source_language
    )

    target = next(
        lang for lang in installed_languages
        if lang.code == target_language
    )

    translator = source.get_translation(target)

    ppt = Presentation(input_file)

    for slide_number, slide in enumerate(ppt.slides, start=1):
        print("\nSlide:", slide_number)

        for shape in slide.shapes:

            # Translate normal text shapes
            if shape.has_text_frame and not shape.has_table:
                for paragraph in shape.text_frame.paragraphs:
                    original_text = paragraph.text

                    if not should_translate(original_text):
                        continue

                    translated_text = translator.translate(original_text)

                    if paragraph.runs:
                        paragraph.runs[0].text = translated_text

                        for run in paragraph.runs[1:]:
                            run.text = ""

                    print(original_text)
                    print("->", translated_text)
                    print("-" * 40)

            # Translate tables while preserving formatting
            if shape.has_table:
                table = shape.table

                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            original_text = paragraph.text

                            if not should_translate(original_text):
                                continue

                            translated_text = translator.translate(
                                original_text
                            )

                            if paragraph.runs:
                                paragraph.runs[0].text = translated_text

                                for run in paragraph.runs[1:]:
                                    run.text = ""

                            print(
                                "Table:",
                                original_text,
                                "->",
                                translated_text
                            )

    ppt.save(output_file)

    print("\nPPT translation completed")
    print("Output:", output_file)

if __name__ == "__main__":
    
    translate_ppt(
      "input/DummyFrench.pptx",
      "output/DummyFrench_English.pptx",
      "fr"
    )